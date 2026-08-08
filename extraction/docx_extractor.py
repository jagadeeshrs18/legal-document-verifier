"""
docx_extractor.py
------------------
Extracts text (paragraphs + tables) from Word documents (.docx).
"""

import docx


def extract_docx(docx_path: str) -> dict:
    document = docx.Document(docx_path)

    paragraphs = [p.text for p in document.paragraphs if p.text.strip()]

    tables = []
    for t_idx, table in enumerate(document.tables):
        rows = []
        for row in table.rows:
            rows.append(" | ".join(cell.text.strip() for cell in row.cells))
        table_text = "\n".join(rows)
        if table_text.strip():
            tables.append({"table_index": t_idx, "text": table_text})

    full_text = "\n\n".join(paragraphs)

    return {
        "file_type": "docx",
        "is_scanned": False,
        "total_paragraphs": len(paragraphs),
        "paragraphs": paragraphs,
        "tables": tables,
        "full_text": full_text,
        "total_chars": len(full_text),
    }
